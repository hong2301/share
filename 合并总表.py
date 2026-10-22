# -*- coding: utf-8 -*-
"""
合并 桌面/20万.xlsx 与 桌面/10万.xlsx -> 桌面/合并总表.xlsx
- 两表表头一致(43 列)
- 按 url 去重(优先保留先读入的 20万 那份)
- 流式读写(openpyxl read_only / write_only), 避免内存爆
"""
from openpyxl import load_workbook, Workbook

SRC = ['/Users/hong/Desktop/20万.xlsx', '/Users/hong/Desktop/10万.xlsx']
OUT = '/Users/hong/Desktop/合并总表.xlsx'

wb_out = Workbook(write_only=True)
ws_out = wb_out.create_sheet('Sheet1')

seen = set()
header_out = None
total_in = 0
dup = 0

for si, path in enumerate(SRC):
    print(f"读取 {path} ...", flush=True)
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    if header_out is None:
        header_out = header
        ws_out.append(header_out)
        print(f"表头({len(header)} 列)已写入", flush=True)
    url_idx = header.index('url') if 'url' in header else None
    n = 0
    for row in it:
        total_in += 1
        key = row[url_idx] if url_idx is not None and url_idx < len(row) else None
        if key is None:
            key = tuple(row[:5])
        if key in seen:
            dup += 1
            continue
        seen.add(key)
        ws_out.append(list(row))
        n += 1
        if n % 20000 == 0:
            print(f"  {path}: 已写入 {n} 行", flush=True)
    wb.close()
    print(f"{path}: 读入 {n} 行(去重后)", flush=True)

print(f"总计读入 {total_in}, 去重 {dup}, 最终 {len(seen)} 行", flush=True)
print("保存中...", flush=True)
wb_out.save(OUT)
print(f"完成: {OUT}", flush=True)